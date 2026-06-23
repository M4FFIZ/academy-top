from pathlib import Path
import subprocess
import sys

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-pptx"])
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt


OUT = Path(r"c:\Users\hahao\Downloads\Telegram Desktop\Презентация_ВКР_Академия_ТОП.pptx")
ASSETS = Path(r"C:\Users\hahao\.cursor\projects\c-Users-hahao-OneDrive-Desktop-academiya\assets")

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

BG = RGBColor(248, 250, 252)
DARK = RGBColor(15, 23, 42)
BLUE = RGBColor(37, 99, 235)
ORANGE = RGBColor(249, 115, 22)
GRAY = RGBColor(71, 85, 105)
WHITE = RGBColor(255, 255, 255)


def set_bg(slide, color=BG):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_accent(slide):
    shape = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(0.16))
    shape.fill.solid()
    shape.fill.fore_color.rgb = ORANGE
    shape.line.fill.background()
    shape = slide.shapes.add_shape(1, Inches(0), Inches(0.16), Inches(13.333), Inches(0.06))
    shape.fill.solid()
    shape.fill.fore_color.rgb = BLUE
    shape.line.fill.background()


def textbox(slide, x, y, w, h, text, size=24, bold=False, color=DARK, align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    if align is not None:
        p.alignment = align
    run = p.runs[0]
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Arial"
    return box


def title(slide, text):
    add_accent(slide)
    textbox(slide, 0.55, 0.42, 12.2, 0.55, text, size=28, bold=True)


def bullets(slide, items, x=0.85, y=1.35, w=11.8, h=5.4, size=22):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = Pt(size)
        p.font.name = "Arial"
        p.font.color.rgb = DARK
        p.space_after = Pt(8)
    return box


def footer(slide, n):
    textbox(slide, 11.95, 7.05, 0.8, 0.25, str(n), size=10, color=GRAY, align=PP_ALIGN.RIGHT)


def add_image(slide, path, x, y, w=None, h=None):
    p = Path(path)
    if p.exists():
        if w and h:
            slide.shapes.add_picture(str(p), Inches(x), Inches(y), width=Inches(w), height=Inches(h))
        elif w:
            slide.shapes.add_picture(str(p), Inches(x), Inches(y), width=Inches(w))
        elif h:
            slide.shapes.add_picture(str(p), Inches(x), Inches(y), height=Inches(h))


def slide_title(subtitle=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, DARK)
    bar = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(0.22))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ORANGE
    bar.line.fill.background()
    textbox(slide, 0.8, 1.15, 11.9, 0.6, "Выпускная квалификационная работа", size=22, color=WHITE)
    textbox(
        slide,
        0.8,
        2.0,
        11.8,
        1.45,
        "Разработка модуля веб-сайта для\nАНО ПОО ММКЦТ «Академия ТОП»",
        size=36,
        bold=True,
        color=WHITE,
    )
    textbox(slide, 0.8, 4.0, 11.6, 0.6, "Модуль личного кабинета образовательной организации", size=22, color=RGBColor(203, 213, 225))
    textbox(slide, 0.8, 6.35, 11.6, 0.35, "Студент: Филишов Михаил", size=18, color=WHITE)
    return slide


slides = []
slides.append(slide_title())

s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); title(s, "Актуальность темы")
bullets(s, [
    "Образовательные организации нуждаются в едином цифровом пространстве для студентов, преподавателей и администрации.",
    "Расписание, оценки, посещаемость и сообщения часто находятся в разных источниках, что снижает удобство работы.",
    "Веб-модуль личного кабинета позволяет централизовать учебную информацию и повысить прозрачность образовательного процесса.",
])
slides.append(s)

s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); title(s, "Цель и задачи работы")
bullets(s, [
    "Цель: разработать модуль веб-сайта для АНО ПОО ММКЦТ «Академия ТОП».",
    "Проанализировать предметную область и аналогичные решения.",
    "Спроектировать архитектуру, базу данных и пользовательские сценарии.",
    "Реализовать frontend, backend, базу данных, чат и административные функции.",
    "Разместить проект на хостинге и провести тестирование.",
])
slides.append(s)

s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); title(s, "Пользователи системы")
bullets(s, [
    "Студент: просмотр расписания, оценок, посещаемости, уведомлений и сообщений.",
    "Преподаватель: ведение электронного журнала, выставление оценок, отметка посещаемости.",
    "Администратор: управление пользователями, группами, дисциплинами, аудиториями и расписанием.",
])
slides.append(s)

s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); title(s, "Функциональные возможности")
bullets(s, [
    "Авторизация и разграничение доступа по ролям.",
    "Личный кабинет пользователя.",
    "Расписание занятий, оценки и посещаемость.",
    "Электронный журнал преподавателя.",
    "Чат, личные сообщения, друзья и уведомления в реальном времени.",
    "Административная панель и аналитика.",
])
slides.append(s)

s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); title(s, "Технологический стек")
bullets(s, [
    "Frontend: Next.js, React, TypeScript, Tailwind CSS, TanStack Query.",
    "Backend: Node.js, NestJS, TypeScript, REST API.",
    "База данных: PostgreSQL, Prisma ORM.",
    "Безопасность: JWT-авторизация, bcrypt-хеширование паролей.",
    "Real-time: Socket.io / WebSocket.",
    "Развёртывание: Amvera Cloud.",
])
slides.append(s)

s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); title(s, "Архитектура решения")
bullets(s, [
    "Клиентская часть отвечает за отображение интерфейса и отправку запросов.",
    "Серверная часть обрабатывает бизнес-логику, авторизацию и доступ к данным.",
    "PostgreSQL хранит учебные данные, сообщения, уведомления и пользователей.",
    "WebSocket обеспечивает мгновенную доставку сообщений и уведомлений.",
], x=0.65, y=1.2, w=5.0, h=5.4, size=19)
add_image(s, ASSETS / "component_diagram_academy_top.png", 6.0, 1.1, w=6.7)
slides.append(s)

s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); title(s, "Информационная база")
bullets(s, [
    "Основные сущности: пользователи, группы, дисциплины, занятия, оценки, посещаемость.",
    "Для коммуникации используются таблицы чатов, сообщений, уведомлений и дружеских связей.",
    "Связи 1:N отражают отношения «один ко многим»: группа — студенты, чат — сообщения, занятие — оценки.",
], x=0.65, y=1.0, w=4.6, h=5.8, size=18)
add_image(s, ASSETS / "er_diagram_academy_top.png", 5.35, 0.95, w=7.5)
slides.append(s)

s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); title(s, "Клиентская часть")
bullets(s, [
    "Реализованы страницы авторизации, личного кабинета, расписания, оценок, посещаемости, чата и профиля.",
    "Интерфейс адаптирован для разных ролей пользователей.",
    "Tailwind CSS обеспечивает единый визуальный стиль и адаптивность.",
    "TanStack Query используется для загрузки и обновления серверных данных.",
])
slides.append(s)

s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); title(s, "Серверная часть")
bullets(s, [
    "NestJS-приложение разделено на функциональные модули.",
    "REST API используется для авторизации, расписания, оценок, посещаемости и администрирования.",
    "Prisma ORM обеспечивает типизированную работу с PostgreSQL.",
    "Socket.io реализует чат и уведомления в реальном времени.",
    "Доступ к закрытым маршрутам защищён JWT-токенами и проверкой ролей.",
])
slides.append(s)

s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); title(s, "Размещение на хостинге")
bullets(s, [
    "Проект размещён в облачной среде Amvera Cloud.",
    "Созданы три компонента инфраструктуры: PostgreSQL, academy-api и academy-web.",
    "Backend доступен по адресу: https://academy-api-maffiz.amvera.io",
    "Frontend доступен по адресу: https://academy-web-maffiz.amvera.io",
    "Соединение выполняется по HTTPS.",
])
slides.append(s)

s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); title(s, "Тестирование")
bullets(s, [
    "Проверена авторизация пользователей с разными ролями.",
    "Протестированы расписание, оценки, посещаемость и электронный журнал.",
    "Проверена отправка сообщений в чате и работа уведомлений.",
    "Проверено ограничение доступа студента к административной панели.",
    "Проверена доступность веб-ресурса по публичной HTTPS-ссылке.",
])
slides.append(s)

s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); title(s, "Контрольный пример")
bullets(s, [
    "Студент входит в личный кабинет и просматривает учебную информацию.",
    "Преподаватель выбирает занятие, выставляет оценку и отмечает посещаемость.",
    "Администратор управляет пользователями, группами, дисциплинами и расписанием.",
    "Данные сохраняются в PostgreSQL и отображаются в интерфейсе после обновления.",
])
slides.append(s)

s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); title(s, "Результаты работы")
bullets(s, [
    "Разработан модуль веб-сайта для образовательной организации.",
    "Реализованы роли студента, преподавателя и администратора.",
    "Созданы frontend, backend, база данных и механизм real-time коммуникации.",
    "Проект размещён на хостинге и доступен в общем доступе.",
    "Проведённое тестирование подтвердило работоспособность основных функций.",
])
slides.append(s)

s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); title(s, "Перспективы развития")
bullets(s, [
    "Добавление загрузки файлов в чат.",
    "Экспорт отчётов и оценок в PDF.",
    "Интеграция с основным сайтом образовательной организации.",
    "Push-уведомления и мобильная версия.",
    "Расширенная аналитика для администрации.",
])
slides.append(s)

s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s, DARK)
textbox(s, 0.8, 2.55, 11.8, 0.8, "Спасибо за внимание!", size=42, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
textbox(s, 0.8, 3.55, 11.8, 0.45, "Готов ответить на вопросы", size=24, color=RGBColor(203, 213, 225), align=PP_ALIGN.CENTER)
slides.append(s)

for i, slide in enumerate(slides, 1):
    if i not in (1, len(slides)):
        footer(slide, i)

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print(OUT)
