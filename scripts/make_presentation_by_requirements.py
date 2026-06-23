from pathlib import Path
import subprocess
import sys

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-pptx"])
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt


OUT = Path(r"C:\Users\hahao\OneDrive\Desktop\Презентация_ВКР_Академия_ТОП_по_требованиям.pptx")
ASSETS = Path(r"C:\Users\hahao\.cursor\projects\c-Users-hahao-OneDrive-Desktop-academiya\assets")

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

WHITE = RGBColor(255, 255, 255)
DARK = RGBColor(20, 30, 45)
BLUE = RGBColor(31, 88, 184)
ORANGE = RGBColor(239, 105, 35)
GRAY = RGBColor(90, 100, 115)
LIGHT = RGBColor(244, 247, 251)


def set_bg(slide, color=WHITE):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def txt(slide, x, y, w, h, text, size=22, bold=False, color=DARK, align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    if align is not None:
        p.alignment = align
    for r in p.runs:
        r.font.name = "Arial"
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
    return box


def title(slide, text):
    band = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(0.18))
    band.fill.solid()
    band.fill.fore_color.rgb = ORANGE
    band.line.fill.background()
    txt(slide, 0.65, 0.45, 12, 0.55, text, size=30, bold=True)


def bullets(slide, items, x=0.9, y=1.45, w=11.7, h=5.2, size=22):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = "Arial"
        p.font.size = Pt(size)
        p.font.color.rgb = DARK
        p.space_after = Pt(12)
    return box


def footer(slide, n):
    txt(slide, 11.85, 7.05, 0.8, 0.25, str(n), size=11, color=GRAY, align=PP_ALIGN.RIGHT)


def image(slide, name, x, y, w=None, h=None):
    path = ASSETS / name
    if path.exists():
        if w and h:
            slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))
        elif w:
            slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w))
        elif h:
            slide.shapes.add_picture(str(path), Inches(x), Inches(y), height=Inches(h))


slides = []

# 1
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s, DARK)
bar = s.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(0.24))
bar.fill.solid()
bar.fill.fore_color.rgb = ORANGE
bar.line.fill.background()
txt(s, 0.85, 1.05, 11.7, 0.45, "Выпускная квалификационная работа", size=22, color=WHITE)
txt(s, 0.85, 2.0, 11.6, 1.35, "Разработка модуля веб-сайта\nдля АНО ПОО ММКЦТ «Академия ТОП»", size=34, bold=True, color=WHITE)
txt(s, 0.85, 4.0, 11.6, 0.45, "Модуль личного кабинета образовательной организации", size=22, color=RGBColor(220, 230, 240))
txt(s, 0.85, 6.35, 11.6, 0.35, "Студент: Филишов Михаил", size=18, color=WHITE)
slides.append(s)

# 2
s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); title(s, "Актуальность")
bullets(s, [
    "Учебная информация часто хранится в разных местах.",
    "Студентам нужно быстро видеть расписание, оценки и сообщения.",
    "Преподавателям нужен удобный электронный журнал.",
    "Администрации важно управлять данными из одного места.",
])
slides.append(s)

# 3
s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); title(s, "Цель и задачи")
bullets(s, [
    "Цель: создать веб-модуль личного кабинета для Академии ТОП.",
    "Изучить предметную область и аналоги.",
    "Спроектировать архитектуру и базу данных.",
    "Реализовать frontend, backend и роли пользователей.",
    "Разместить проект на хостинге и протестировать.",
])
slides.append(s)

# 4
s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); title(s, "Роли пользователей")
bullets(s, [
    "Студент: расписание, оценки, посещаемость, чат.",
    "Преподаватель: журнал, оценки, посещаемость, сообщения.",
    "Администратор: пользователи, группы, дисциплины, расписание.",
])
slides.append(s)

# 5
s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); title(s, "Основные функции")
bullets(s, [
    "Авторизация и защита личного кабинета.",
    "Расписание занятий.",
    "Оценки и посещаемость.",
    "Электронный журнал преподавателя.",
    "Чат и уведомления.",
    "Административная панель.",
])
slides.append(s)

# 6
s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); title(s, "Технологии")
bullets(s, [
    "Frontend: Next.js, React, TypeScript, Tailwind CSS.",
    "Backend: Node.js, NestJS, REST API.",
    "База данных: PostgreSQL и Prisma ORM.",
    "Авторизация: JWT и bcrypt.",
    "Чат: Socket.io / WebSocket.",
    "Хостинг: Amvera Cloud.",
])
slides.append(s)

# 7
s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); title(s, "Архитектура проекта")
bullets(s, [
    "Frontend показывает интерфейс пользователю.",
    "Backend обрабатывает запросы и права доступа.",
    "PostgreSQL хранит учебные данные.",
    "WebSocket нужен для чата и уведомлений.",
], x=0.65, y=1.25, w=4.3, h=5.1, size=20)
image(s, "component_diagram_academy_top.png", 5.15, 1.05, w=7.65)
slides.append(s)

# 8
s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); title(s, "База данных")
bullets(s, [
    "Хранятся пользователи, группы, предметы и занятия.",
    "Оценки и посещаемость связаны со студентами.",
    "Чаты, сообщения и уведомления хранятся отдельно.",
], x=0.65, y=1.25, w=4.2, h=5.1, size=20)
image(s, "er_diagram_academy_top.png", 5.0, 1.0, w=7.85)
slides.append(s)

# 9
s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); title(s, "Реализация и размещение")
bullets(s, [
    "Созданы два сервиса: academy-web и academy-api.",
    "Отдельно создана база PostgreSQL.",
    "Проект размещён на Amvera.",
    "Сайт доступен по HTTPS-ссылке.",
    "API работает отдельно и отдаёт данные сайту.",
])
slides.append(s)

# 10
s = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(s); title(s, "Тестирование и результат")
bullets(s, [
    "Проверена авторизация для всех ролей.",
    "Проверены расписание, оценки и посещаемость.",
    "Проверены журнал преподавателя, чат и админ-панель.",
    "Проект работает на хостинге и готов к демонстрации.",
    "Цель работы достигнута.",
])
slides.append(s)

# 11
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s, DARK)
txt(s, 0.8, 2.6, 11.7, 0.7, "Спасибо за внимание!", size=42, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
txt(s, 0.8, 3.55, 11.7, 0.4, "Готов ответить на вопросы", size=24, color=RGBColor(220, 230, 240), align=PP_ALIGN.CENTER)
slides.append(s)

for i, s in enumerate(slides, 1):
    if i not in (1, len(slides)):
        footer(s, i)

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print(OUT)
