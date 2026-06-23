from pathlib import Path
import re
import shutil
import subprocess
import sys

try:
    from docx import Document
    from docx.enum.section import WD_SECTION_START
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    from docx import Document
    from docx.enum.section import WD_SECTION_START
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor


SRC = Path(r"c:\Users\hahao\OneDrive\Desktop\Диплом Евтушенко.docx")
BACKUP = SRC.with_name("Диплом Евтушенко_до_оформления.docx")
OUT = SRC.with_name("Диплом Евтушенко_ГОСТ.docx")


def set_run_font(run, size=14, bold=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.font.bold = bold


def set_paragraph_base(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Cm(1.25)
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    for run in paragraph.runs:
        set_run_font(run, 14)


def is_main_heading(text: str) -> bool:
    t = text.strip()
    up = t.upper()
    if up in {
        "ВВЕДЕНИЕ",
        "ЗАКЛЮЧЕНИЕ",
        "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ",
        "СПИСОК ЛИТЕРАТУРЫ",
        "СОДЕРЖАНИЕ",
        "ОГЛАВЛЕНИЕ",
    }:
        return True
    if re.match(r"^(РАЗДЕЛ\s+\d+|ГЛАВА\s+\d+)(\s+|$)", up):
        return True
    return False


def is_subheading(text: str) -> bool:
    t = text.strip()
    return bool(re.match(r"^\d+\.\d+(\.\d+)?\s+\S+", t))


def is_caption(text: str) -> bool:
    return bool(re.match(r"^(Рисунок|Таблица)\s+\d+(\.\d+)?\s*[—-]", text.strip(), re.I))


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    set_run_font(run, 12)


def format_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.first_line_indent = Cm(0)
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    set_run_font(run, 12)


def main():
    if not SRC.exists():
        raise FileNotFoundError(SRC)

    if not BACKUP.exists():
        shutil.copy2(SRC, BACKUP)

    doc = Document(str(SRC))

    # Page setup: A4 margins.
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(3)
        section.right_margin = Cm(1.5)
        section.different_first_page_header_footer = True

    # Base styles.
    for style_name in ["Normal", "Body Text"]:
        if style_name in doc.styles:
            style = doc.styles[style_name]
            style.font.name = "Times New Roman"
            style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
            style.font.size = Pt(14)
            style.font.color.rgb = RGBColor(0, 0, 0)

    seen_first_main = False
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        set_paragraph_base(paragraph)

        if not text:
            paragraph.paragraph_format.first_line_indent = Cm(0)
            continue

        if is_main_heading(text):
            # Do not force page break before contents/title-like first heading.
            if seen_first_main and text.upper() not in {"СОДЕРЖАНИЕ", "ОГЛАВЛЕНИЕ"}:
                paragraph.paragraph_format.page_break_before = True
            seen_first_main = True
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.line_spacing = 1.5
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(12)
            if text.upper() != text and not re.match(r"^\d", text):
                paragraph.text = text.upper()
            for run in paragraph.runs:
                set_run_font(run, 14, bold=True)

        elif is_subheading(text):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.first_line_indent = Cm(1.25)
            paragraph.paragraph_format.space_before = Pt(12)
            paragraph.paragraph_format.space_after = Pt(6)
            for run in paragraph.runs:
                set_run_font(run, 14, bold=True)

        elif is_caption(text):
            # Figures centered, tables left.
            if text.lower().startswith("рисунок"):
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(6)
            for run in paragraph.runs:
                set_run_font(run, 14)

    for table in doc.tables:
        format_table(table)

    # Add page numbers in footer for all sections. First page is hidden by different_first_page.
    for section in doc.sections:
        footer = section.footer
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.clear()
        add_page_number(p)

    doc.save(str(OUT))
    print(f"Saved: {OUT}")
    print(f"Backup: {BACKUP}")


if __name__ == "__main__":
    main()
